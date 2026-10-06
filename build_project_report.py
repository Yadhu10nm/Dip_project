import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for table cells."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None, color="D3D3D3", sz="4"):
    """Set light borders for table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name, active in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if active:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def add_styled_heading(doc, text, level=1, space_before=12, space_after=6):
    """Add a heading with strict Times New Roman formatting."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 2:
        run.font.size = Pt(14)
        run.bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 3:
        run.font.size = Pt(12)
        run.bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def add_body_p(doc, text="", space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False):
    """Add standard justified body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, bold_prefix, text, space_after=4):
    """Add bullet paragraph with bold prefix and left indent."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.size = Pt(12)
    r_bullet.bold = True
    r_bullet.font.color.rgb = RGBColor(0, 0, 0)
    
    r_bold = p.add_run(bold_prefix + ": ")
    r_bold.font.name = 'Times New Roman'
    r_bold.font.size = Pt(12)
    r_bold.bold = True
    r_bold.font.color.rgb = RGBColor(0, 0, 0)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)
    r_text.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_block(doc, code_text):
    """Add formatted monospace code snippet with light gray shading."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, 'F3F4F6')
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_borders(cell, top=True, bottom=True, left=True, right=True, color='9CA3AF', sz='6')
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(17, 24, 39)
    
    # spacing paragraph after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, data, col_widths, caption=""):
    """Add a professional academic table with shaded headers and thin borders."""
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11)
        c_run.bold = True
        c_run.font.color.rgb = RGBColor(0, 0, 0)
        
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], '1E3A8A')
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_borders(hdr_cells[i], top=True, bottom=True, left=True, right=True, color='1E3A8A', sz='6')
        hdr_cells[i].width = col_widths[i]
        
        hp = hdr_cells[i].paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hp.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = 'F9FAFB' if r_idx % 2 == 1 else 'FFFFFF'
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=140, right=140)
            set_cell_borders(row_cells[c_idx], top=True, bottom=True, left=True, right=True, color='D1D5DB', sz='4')
            row_cells[c_idx].width = col_widths[c_idx]
            
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 and len(val) < 8 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0, 0, 0)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def add_figure_with_caption(doc, image_path, caption, width_inches=5.8):
    """Add a centered high-resolution figure with academic caption."""
    if os.path.exists(image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(10.5)
        c_run.italic = True
        c_run.font.color.rgb = RGBColor(55, 65, 81)
    else:
        print(f"Warning: Image not found: {image_path}")

print("Helper definitions loaded successfully.")
