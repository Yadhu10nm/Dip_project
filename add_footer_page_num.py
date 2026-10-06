import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

def update_footer(docx_path):
    doc = docx.Document(docx_path)
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ''
    
    run = p.add_run("Surana College Autonomous | Department of Computer Science | Page ")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    r2 = p.add_run()
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(107, 114, 128)
    
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    r2._r.append(fld3)
    
    doc.save(docx_path)
    print("Dynamic footer page number updated successfully.")

if __name__ == "__main__":
    update_footer("project_report.docx")
